from pypdf import PdfReader
from docx import Document


class UnsupportedFileType(Exception):
    """Raised when the file extension has no registered reader."""


class ExtractionFailed(Exception):
    """Raised when a supported file exists but cannot be read (corrupt, encrypted)."""


def extract_text(file_path: str) -> str:
    """Read a supported document and return its full text.

    Images, charts and other non-text content are skipped, not fatal.
    Unsupported types raise UnsupportedFileType; unreadable files raise
    ExtractionFailed. Neither crashes the service.
    """
    lower = file_path.lower()

    if lower.endswith(".pdf"):
        try:
            reader = PdfReader(file_path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise ExtractionFailed(f"Could not read PDF: {e}") from e

    if lower.endswith(".docx"):
        try:
            doc = Document(file_path)
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(part for part in parts if part.strip())
        except Exception as e:
            raise ExtractionFailed(f"Could not read DOCX: {e}") from e

    if lower.endswith((".txt", ".md")):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            raise ExtractionFailed(f"Could not read text file: {e}") from e

    raise UnsupportedFileType(
        f"Unsupported file type: {file_path.rsplit('.', 1)[-1]}. "
        "Supported: .pdf, .docx, .txt, .md"
    )


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    """Split text into overlapping chunks of roughly `chunk_size` words."""
    words = text.split()
    if not words:
        return []

    chunks = []
    step = chunk_size - overlap
    for start in range(0, len(words), step):
        chunk = words[start:start + chunk_size]
        if chunk:
            chunks.append(" ".join(chunk))
        if start + chunk_size >= len(words):
            break
    return chunks